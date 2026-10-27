"""Belge yükleyiciler (PDF, DOCX, TXT/MD, web) ve metin parçalama."""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

log = logging.getLogger(__name__)
Progress = Callable[[int, int], None] | None

SUPPORTED_EXT = {".pdf", ".docx", ".txt", ".md", ".markdown"}

CHUNK_SIZE = 1200      # karakter
CHUNK_OVERLAP = 200    # karakter


@dataclass
class LoadedDoc:
    title: str
    source: str
    kind: str
    sections: list[tuple[str, str]]  # (konum etiketi, metin)
    warnings: list[str] = field(default_factory=list)


class IngestError(Exception):
    pass


# ---------------------------------------------------------------- loaders
def _pdf_pages_pdfium(path: Path, progress: Progress):
    """pypdfium2 ile hızlı ve Türkçe karakterlerde daha güvenilir metin çıkarımı."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    try:
        n = len(pdf)
        title = ""
        try:
            title = (pdf.get_metadata_dict() or {}).get("Title", "") or ""
        except Exception:  # noqa: BLE001
            pass
        for i in range(n):
            try:
                page = pdf[i]
                tp = page.get_textpage()
                txt = tp.get_text_range() or ""
                tp.close()
                page.close()
            except Exception:  # noqa: BLE001
                txt = ""
            if progress and (i % 10 == 0 or i == n - 1):
                progress(i + 1, n)
            yield i + 1, n, txt.replace("\r\n", "\n").replace("\r", "\n"), title
    finally:
        pdf.close()


def _pdf_pages_pypdf(path: Path, progress: Progress):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        reader.decrypt("")
    n = len(reader.pages)
    title = (reader.metadata.title if reader.metadata and reader.metadata.title else "") or ""
    for i, page in enumerate(reader.pages):
        try:
            txt = page.extract_text() or ""
        except Exception:  # noqa: BLE001
            txt = ""
        if progress and (i % 10 == 0 or i == n - 1):
            progress(i + 1, n)
        yield i + 1, n, txt, title


def load_pdf(path: Path, progress: Progress = None) -> LoadedDoc:
    sections: list[tuple[str, str]] = []
    title, n_pages, empty = "", 0, 0
    readers = []
    try:
        import pypdfium2  # noqa: F401
        readers.append(_pdf_pages_pdfium)
    except ImportError:
        pass
    readers.append(_pdf_pages_pypdf)
    last_err: Exception | None = None
    for reader in readers:
        sections, empty, n_pages = [], 0, 0
        try:
            for i, n, txt, t in reader(path, progress):
                n_pages, title = n, t or title
                if len(txt.strip()) < 20:
                    empty += 1
                    continue
                sections.append((f"s. {i}", txt))
            last_err = None
            break
        except Exception as e:  # noqa: BLE001
            last_err = e
            log.warning("PDF okuyucu %s başarısız: %s", reader.__name__, e)
    if last_err is not None:
        msg = str(last_err).lower()
        if "password" in msg or "encrypt" in msg:
            raise IngestError("PDF şifreli, açılamadı.") from last_err
        raise IngestError(f"PDF açılamadı: {last_err}") from last_err
    if not sections:
        raise IngestError(
            "PDF'ten metin çıkarılamadı. Taranmış (resim) bir PDF olabilir; "
            "önce OCR ile metne çevirmeniz gerekir."
        )
    warnings = []
    if n_pages and empty / n_pages > 0.3:
        warnings.append(f"{n_pages} sayfanın {empty} tanesinde metin bulunamadı (taranmış sayfalar olabilir); "
                        "bu sayfalardan soru üretilemez.")
    sample = " ".join(t for _, t in sections[:: max(1, len(sections) // 20)])[:20000]
    if sample:
        nonspace = [c for c in sample if not c.isspace()]
        letters = sum(c.isalpha() for c in nonspace) / max(1, len(nonspace))
        if letters < 0.6:
            warnings.append("Metin çıkarımı bozuk görünüyor (çok sayıda anlamsız karakter). Kitap düzgün okunmamış "
                            "olabilir; mümkünse farklı bir PDF sürümü deneyin.")
    if str(title).strip().lower() in ("", "untitled", "microsoft word", "document") or len(str(title)) < 3:
        title = path.stem
    return LoadedDoc(title=str(title).strip(), source=str(path), kind="pdf", sections=sections, warnings=warnings)


def load_docx(path: Path) -> LoadedDoc:
    import docx

    try:
        d = docx.Document(str(path))
    except Exception as e:  # noqa: BLE001
        raise IngestError(f"Word dosyası açılamadı: {e}") from e
    sections: list[tuple[str, str]] = []
    current_head = ""
    buf: list[str] = []

    def flush():
        if buf:
            sections.append((current_head, "\n".join(buf)))
            buf.clear()

    for p in d.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") or style.startswith("başlık"):
            flush()
            current_head = text[:80]
            buf.append(text)
        else:
            buf.append(text)
    flush()
    # tablolar
    for ti, table in enumerate(d.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            sections.append((f"Tablo {ti}", "\n".join(rows)))
    if not sections:
        raise IngestError("Word dosyasında metin bulunamadı.")
    title = d.core_properties.title or path.stem
    return LoadedDoc(title=title, source=str(path), kind="docx", sections=sections)


def load_text(path: Path) -> LoadedDoc:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp1254", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if not text.strip():
        raise IngestError("Dosya boş.")
    return LoadedDoc(title=path.stem, source=str(path), kind=path.suffix.lstrip(".").lower(), sections=[("", text)])


def load_url(url: str, timeout: int = 25) -> LoadedDoc:
    import requests
    from bs4 import BeautifulSoup

    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) MedQuiz/1.0"}
    try:
        resp = requests.get(url, headers=headers, timeout=timeout)
        resp.raise_for_status()
    except Exception as e:  # noqa: BLE001
        raise IngestError(f"Sayfa indirilemedi: {e}") from e

    ctype = resp.headers.get("content-type", "")
    if "pdf" in ctype or url.lower().endswith(".pdf"):
        import tempfile

        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
            f.write(resp.content)
            tmp = Path(f.name)
        try:
            doc = load_pdf(tmp)
        finally:
            tmp.unlink(missing_ok=True)
        doc.source, doc.kind = url, "web-pdf"
        return doc

    resp.encoding = resp.apparent_encoding or resp.encoding
    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "aside", "form", "iframe", "svg"]):
        tag.decompose()
    title = (soup.title.string.strip() if soup.title and soup.title.string else url)[:200]
    main = soup.find("article") or soup.find("main") or soup.body or soup
    sections: list[tuple[str, str]] = []
    head, buf = "", []
    for el in main.find_all(["h1", "h2", "h3", "h4", "p", "li", "td", "th", "pre", "blockquote", "dd", "dt"]):
        txt = el.get_text(" ", strip=True)
        if not txt:
            continue
        if el.name in ("h1", "h2", "h3", "h4"):
            if buf:
                sections.append((head, "\n".join(buf)))
                buf = []
            head = txt[:80]
            buf.append(txt)
        elif len(txt) > 25 or el.name in ("li", "td", "dd"):
            buf.append(txt)
    if buf:
        sections.append((head, "\n".join(buf)))
    if sum(len(s[1]) for s in sections) < 200:
        text = main.get_text("\n", strip=True)
        if len(text) < 200:
            raise IngestError("Sayfada yeterli metin bulunamadı (giriş gerektiren ya da JavaScript ile yüklenen bir sayfa olabilir).")
        sections = [("", text)]
    return LoadedDoc(title=title, source=url, kind="web", sections=sections)


def load_any(source: str, progress: Progress = None) -> LoadedDoc:
    if re.match(r"^(https?://|www\.)", source.strip(), re.I):
        return load_url(source.strip())
    path = Path(source)
    if not path.exists():
        raise IngestError(f"Dosya bulunamadı: {source}")
    ext = path.suffix.lower()
    if ext == ".pdf":
        return load_pdf(path, progress)
    if ext == ".docx":
        return load_docx(path)
    if ext in (".txt", ".md", ".markdown"):
        return load_text(path)
    if ext == ".doc":
        raise IngestError("Eski .doc biçimi desteklenmiyor; Word'de 'Farklı Kaydet → .docx' yapın.")
    raise IngestError(f"Desteklenmeyen dosya türü: {ext}")


# --------------------------------------------------------------- chunking
_SENT_SPLIT = re.compile(r"(?<=[.!?…:;])\s+(?=[A-ZÇĞİÖŞÜ0-9•\-–(])")


def _clean(text: str) -> str:
    text = text.replace("\r", "\n").replace("­", "")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)          # satır sonu tirelemesi
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split_units(text: str) -> list[str]:
    """Paragraf → cümle → kelime sırasıyla küçük birimlere böler."""
    units: list[str] = []
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        if len(para) <= CHUNK_SIZE:
            units.append(para)
            continue
        for sent in _SENT_SPLIT.split(para):
            sent = sent.strip()
            if len(sent) <= CHUNK_SIZE:
                units.append(sent)
            else:  # çok uzun "cümle" (tablo vb.)
                words = sent.split()
                cur = ""
                for w in words:
                    if len(cur) + len(w) + 1 > CHUNK_SIZE:
                        units.append(cur)
                        cur = w
                    else:
                        cur = f"{cur} {w}".strip()
                if cur:
                    units.append(cur)
    return units


def chunk_doc(doc: LoadedDoc) -> list[dict]:
    chunks: list[dict] = []
    for loc, raw in doc.sections:
        text = _clean(raw)
        if not text:
            continue
        units = _split_units(text)
        cur: list[str] = []
        cur_len = 0
        for u in units:
            if cur and cur_len + len(u) + 1 > CHUNK_SIZE:
                chunks.append({"text": "\n".join(cur), "location": loc})
                # örtüşme: sondan geriye CHUNK_OVERLAP kadar birim taşı
                keep, keep_len = [], 0
                for prev in reversed(cur):
                    if keep_len + len(prev) > CHUNK_OVERLAP:
                        break
                    keep.insert(0, prev)
                    keep_len += len(prev) + 1
                cur, cur_len = keep, keep_len
            cur.append(u)
            cur_len += len(u) + 1
        if cur:
            chunks.append({"text": "\n".join(cur), "location": loc})

    # çok küçük parçaları bir öncekiyle birleştir
    merged: list[dict] = []
    for c in chunks:
        if merged and len(c["text"]) < 200 and len(merged[-1]["text"]) + len(c["text"]) < CHUNK_SIZE * 1.3:
            merged[-1]["text"] += "\n" + c["text"]
        else:
            merged.append(c)
    return merged
